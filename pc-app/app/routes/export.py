from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime
from typing import Optional
from app.database import get_db
from app.models import Attendance

router = APIRouter(prefix="/api/export", tags=["export"])


@router.get("/excel")
async def export_excel(
    start: Optional[str] = Query(None),
    end: Optional[str] = Query(None),
    uid: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    query = db.query(Attendance)

    if uid:
        query = query.filter(Attendance.uid == uid)
    if status:
        query = query.filter(Attendance.status == status)
    if start:
        try:
            query = query.filter(Attendance.timestamp >= datetime.fromisoformat(start))
        except:
            pass
    if end:
        try:
            query = query.filter(Attendance.timestamp <= datetime.fromisoformat(end))
        except:
            pass

    records = query.order_by(Attendance.timestamp).all()

    # Build Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Absensi"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Headers
    headers = ["No", "UID", "Nama", "Tanggal", "Jam", "Status"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data rows
    for i, att in enumerate(records, 1):
        ts_local = att.timestamp.strftime("%Y-%m-%d")
        tm_local = att.timestamp.strftime("%H:%M:%S")
        status_text = "MASUK" if att.status == "IN" else "KELUAR"

        row_data = [i, att.uid, att.name, ts_local, tm_local, status_text]
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=i + 1, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col != 3 else "left")

    # Auto-fit column width
    for col in range(1, 7):
        max_length = len(str(headers[col - 1]))
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=True):
            for cell_value in row:
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
        ws.column_dimensions[chr(64 + col)].width = max_length + 3

    # Save to memory
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    # Generate filename
    filename = f"absensi_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
